"""
tl_journal.py - tl_journal v1.0, an add-on for tradelab v1.3.1: ONE readable Excel file for an export.

    python tl_journal.py --export TLX_20261025_001
    python tl_journal.py --export C:\\QuantData\\tradelab\\export\\TLX_20261025_001
        [--all]          also the picks that were not traded (ranks 4-10): about 3x the rows
        [--name FILE]    the file name inside the export folder (default TRADE_JOURNAL.xlsx)
    python tl_journal.py --selftest       builds a small synthetic export, writes its journal and checks it cell by cell

It reads the export READ ONLY - trades.parquet (checked against the export's outputs.sha256), market_days.parquet and
context_breakdown.csv - and writes one file, <export>\\TRADE_JOURNAL.xlsx. Nothing else in the export changes.

  Trades       one row per traded pick (D and ENS ranks 1-3), newest first, the columns in bands: the trade, its
               result, the close day by day, the trade's details, the stock's state at the pick, the market at the
               pick (environment, breadth, VIX), the hold (market part vs stock part) and the story in one line
  Market days  one row per session, newest first: environment, indices, breadth, VIX, and how that day's picks did
  What worked  tradelab's own breakdown of the traded picks (context_breakdown.csv) with its 95% intervals, copied as
               they are - nothing is recomputed
  Read me      the colours and every column in plain words

Values, not formulas: the journal is a view of the export. Delete it and run this again any time.
Needs openpyxl (tradelab's AMFI reader uses it too): pip install openpyxl
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from tl_common import TLError, VERSION, say      # noqa: E402
import tl_context as C                           # noqa: E402
import tl_state as ST                            # noqa: E402

TOOL = "tl_journal v1.0"
NEEDS = "tradelab v1.3.1"
DEFAULT_NAME = "TRADE_JOURNAL.xlsx"
FONT = "Arial"
ENGINE_ORDER = ("D", "ENS")

# ----------------------------------------------------------------------------------------------------------------
# colours (fill, font) - the Read me sheet shows every one of them
# ----------------------------------------------------------------------------------------------------------------
ZEBRA = "F2F2F2"
GOOD, GOOD_TXT = "C6EFCE", "006100"
BAD, BAD_TXT = "FFC7CE", "9C0006"
RESULT_STYLE = {"Winner": (GOOD, GOOD_TXT, False), "Loser": (BAD, BAD_TXT, False),
                "Severe loser": ("C00000", "FFFFFF", True), "Stuck": ("7F7F7F", "FFFFFF", True)}
ENV_STYLE = {"broad_bullish": ("C6EFCE", "006100", False), "narrow_bullish": ("E2F0D9", "375623", False),
             "mixed": ("E7E6E6", "404040", False), "broad_bearish": ("F8CBAD", "843C0C", False),
             "high_stress": ("7030A0", "FFFFFF", True)}
STATE_STYLE = {"S01": ("A9D08E", "1E4620", False), "S02": ("C6E0B4", "375623", False),
               "S03": ("E2EFDA", "375623", False), "S06": ("DDEBF7", "1F4E79", False),
               "S07": ("E7E6E6", "404040", False), "S08": ("F4B084", "843C0C", False),
               "S09": ("F8CBAD", "843C0C", False), "S12": ("FFE699", "7F6000", False)}
VOL_STYLE = {"EXTREME": ("F4B084", "843C0C", False), "ELEVATED": ("FCE4D6", "843C0C", False),
             "COMPRESSED": ("DDEBF7", "1F4E79", False)}
VERDICT_STYLE = {"above 0": (GOOD, GOOD_TXT, True), "below 0": (BAD, BAD_TXT, True), "spans 0": ("F2F2F2", "595959", False)}
RED, WHITE, GREEN = "F8696B", "FFFFFF", "63BE7B"
SCALES = {   # conditional colour scales: (value, colour) points - fixed, so a colour means the same in every file
    "path": ((-0.08, RED), (0.0, WHITE), (0.08, GREEN)),
    "best": ((0.0, WHITE), (0.15, GREEN)),
    "worst": ((-0.15, RED), (0.0, WHITE)),
    "breadth": ((0.25, RED), (0.5, WHITE), (0.75, GREEN)),
    "vix5": ((-0.30, GREEN), (0.0, WHITE), (0.30, RED)),
    "vixpct": ((0.0, WHITE), (0.5, WHITE), (1.0, RED)),
    "rate": ((0.0, WHITE), (1.0, "9BC2E6")),
}
NF = {"date": "yyyy-mm-dd ddd", "pct2": "+0.00%;-0.00%;0.00%", "pct1": "+0.0%;-0.0%;0.0%", "pct0": "0%",
      "ratio": "0.00", "price": "#,##0.00", "int": "0", "num0": "0", "num1": "0.0", "num2": "0.00",
      "bp": "+#,##0;-#,##0;0", "count": "#,##0", "share": "0.0%", "text": None}

PHASE_TEXT = {"EARLY": "Early", "MATURE": "Mature", "DETERIORATING": "Deteriorating"}
CONF_TEXT = {"HIGH": "High", "MEDIUM": "Medium", "LOW": "Low"}
VOL_TEXT = {"COMPRESSED": "Compressed", "NORMAL": "Normal", "ELEVATED": "Elevated", "EXTREME": "Extreme"}


def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise TLError("openpyxl is not installed: pip install openpyxl") from None


# ----------------------------------------------------------------------------------------------------------------
# small value helpers
# ----------------------------------------------------------------------------------------------------------------
def _f(x) -> Optional[float]:
    """A finite float, else None (blank cell)."""
    return float(x) if C.fin(x) else None


def _i(x) -> Optional[int]:
    return int(round(float(x))) if C.fin(x) else None


def _b(x) -> bool:
    """True only for a real True / 1 (a blank or NaN is False - bool(nan) would be True)."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return False
    try:
        return bool(x) and not pd.isna(x)
    except (TypeError, ValueError):
        return bool(x)


def _px(x) -> str:
    return f"{float(x):,.2f}" if C.fin(x) else "n/a"


def _s(x) -> str:
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return ""
    s = str(x)
    return "" if s in ("nan", "None", "NaT", "<NA>") else s


def _day(x) -> Optional[dt.date]:
    if x is None or (not isinstance(x, (str, dt.date)) and pd.isna(x)):
        return None
    try:
        t = pd.Timestamp(x)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(t) else t.date()


def _div(x, k: float) -> Optional[float]:
    v = _f(x)
    return None if v is None else v / k


def state_label(code: str, full: bool = False) -> str:
    code = _s(code)
    if not code:
        return ""
    if code in ("UNKNOWN", "INSUFFICIENT_DATA"):
        return (ST.STATE_NAMES if full else ST.SHORT).get(code, code)
    nm = (ST.STATE_NAMES if full else ST.SHORT).get(code) or ST.STATE_NAMES.get(code, "")
    return f"{code} {nm}".strip()


EVENT_SHORT = {"BREAKOUT": "breakout", "FAILED_BREAKOUT": "failed breakout", "BREAKDOWN": "breakdown",
               "FAILED_BREAKDOWN": "failed breakdown"}
RISK_SHORT = {"S05_EXTREME": "S05 extreme", "S05_OVEREXTENDED": "S05 overextended", "STRETCHED_DOWN": "stretched down",
              "BIG_DOWN_CANDLE": "big down candle", "GAP_RISK": "gap risk", "BEAR_DIVERGENCE": "bear divergence",
              "PRICE_UP_FLOW_DOWN": "price up, flow down"}
NOTE_SHORT = {"BULL_DIVERGENCE": "bull divergence", "PRICE_DOWN_FLOW_UP": "price down, flow up",
              "VOL_COMPRESSED": "vol compressed"}


def event_text(event: str, age, short: bool = False) -> str:
    """'S04 breakout confirmed (3d)' / 'S11 failed breakout (bull trap) (5d)' / '' (tl_state's names; short = no
    '(bull trap)')."""
    ev = _s(event)
    if not ev:
        return ""
    kind, _, status = ev.partition(":")
    code, nm = ST.EVENT_NAMES.get(kind, ("", kind.lower()))
    if short:
        nm = EVENT_SHORT.get(kind, nm)
    s = f"{code + ' ' if code else ''}{nm}"
    if status and status != "FAILED":
        s += f" {status.lower()}"
    a = _i(age)
    if a is not None:
        s += " today" if a == 0 else f" ({a}d)"
    return s


def flags_text(risks: str, notes: str, short: bool = False) -> str:
    rk = [(RISK_SHORT if short else ST.RISK_TEXT).get(x, x) for x in _s(risks).split(",") if x]
    nt = [(NOTE_SHORT if short else ST.NOTE_TEXT).get(x, x) for x in _s(notes).split(",") if x]
    parts = []
    if rk:
        parts.append("risk: " + "; ".join(rk))
    if nt:
        parts.append("note: " + "; ".join(nt))
    return " | ".join(parts)


def result_of(r: dict) -> str:
    """Winner / Loser / Severe loser (tradelab's rule: net > 0 / 0 to -500 bp / -500 bp or worse) / Stuck."""
    if _s(r.get("status")) != "ok" or not C.fin(r.get("net_bp")):
        return "Stuck" if _s(r.get("status")) == "stuck" else (_s(r.get("status")).capitalize() or "No result")
    v = float(r["net_bp"])
    return "Winner" if v > 0 else ("Loser" if v > C.SEVERE_BP else "Severe loser")


def exit_text(r: dict) -> str:
    k = _s(r.get("exit_kind"))
    n = _i(r.get("sessions_held"))
    if k == "close":
        return f"day-{n} close" if n else "close"
    if k == "locked_open":
        return f"locked: sold day {n} open" if n else "locked: sold at the next free open"
    if k == "stuck_close":
        return "stuck: locked 20+ sessions"
    return k


def notes_text(r: dict, other_engine: Dict[str, str]) -> str:
    out = []
    oth = other_engine.get(_s(r.get("engine")), "the other engine")
    if "traded" in r and not _b(r.get("traded")):
        out.append(f"not traded (rank {_i(r.get('rank'))})")
    if _b(r.get("both_engines")):
        k = _i(r.get("other_rank"))
        out.append(f"also {oth} #{k}" if k else f"also traded by {oth}")
    if _b(r.get("repeat")) and _s(r.get("repeat_of")):
        out.append(f"repeat of {_s(r.get('repeat_of'))}")
    if _s(r.get("note")):
        out.append(_s(r.get("note")))
    return "; ".join(out)


def story(r: dict) -> str:
    """The pick in one line: the market that day -> the stock and its state -> the trade -> the hold."""
    g = r.get
    env = C.ENV_TEXT.get(_s(g("pre_environment")), "unknown")
    s1 = f"{C.day_label(g('signal_day'))}: {env} market"
    if _s(g("pre_environment")) == "high_stress" and _s(g("pre_stress_reason")):
        s1 += f" ({_s(g('pre_stress_reason'))})"
    bits = [f"NIFTY {C.pct(g('pre_nifty_ret'))}", f"Smallcap {C.pct(g('pre_small_ret'))}"]
    if C.fin(g("pre_ad_ratio")):
        b = f"breadth {float(g('pre_ad_ratio')):.2f}"
        if C.fin(g("pre_ad_ratio_10")):
            b += f" (10-day {float(g('pre_ad_ratio_10')):.2f})"
        bits.append(b)
    if C.fin(g("pre_vix")):
        bits.append(f"VIX {float(g('pre_vix')):.1f} ({C.pct(g('pre_vix_chg5_pct'), 0)} in 5 days)")
    s1 += " - " + ", ".join(bits) + "."
    st = _s(g("pre_state"))
    rk = _i(g("rank"))
    s2 = f"{_s(g('engine'))} #{rk if rk is not None else '?'} {_s(g('symbol'))}"
    if st:
        if st in ("UNKNOWN", "INSUFFICIENT_DATA"):
            s2 += f": state {ST.STATE_NAMES.get(st, st)}"
            if C.fin(g("pre_state_age")):
                s2 += f" ({_i(g('pre_state_age'))} sessions)"
        else:
            det = [x for x in (CONF_TEXT.get(_s(g("pre_state_confidence")), "").lower() and
                               f"{CONF_TEXT[_s(g('pre_state_confidence'))].lower()} confidence",
                               C.fin(g("pre_state_age")) and f"{_i(g('pre_state_age'))} sessions",
                               PHASE_TEXT.get(_s(g("pre_state_phase")), "").lower()) if x]
            s2 += f": {state_label(st, full=True)}" + (f" ({', '.join(det)})" if det else "")
        if _s(g("pre_state_prev")):
            s2 += f", was {state_label(g('pre_state_prev'))}"
        if _s(g("pre_state_pending")):
            s2 += f", {state_label(g('pre_state_pending'))} pending"
        ev = event_text(g("pre_state_event"), g("pre_state_event_age"))
        if ev:
            s2 += f"; {ev}"
        fl = flags_text(g("pre_state_risks"), g("pre_state_notes"))
        if fl:
            s2 += "; " + fl.replace(" | ", "; ")
        if _s(g("pre_state_vol_regime")) in ("ELEVATED", "EXTREME", "COMPRESSED"):
            s2 += f"; volatility {_s(g('pre_state_vol_regime')).lower()}"
    if _s(g("pre_home_index_name")):
        s2 += f"; moves with {_s(g('pre_home_index_name'))}"
        if C.fin(g("pre_beta_home")):
            s2 += f" (beta {float(g('pre_beta_home')):.2f})"
    s2 += "."
    res = result_of(r)
    if res in ("Winner", "Loser", "Severe loser"):
        s3 = f"Bought {C.day_label(g('entry_day'))} open at {_px(g('entry'))}"
        if C.fin(g("entry_gap")):
            s3 += f" (gap {C.pct(g('entry_gap'))})"
        s3 += (f"; best {C.pct(g('mfe'))} on day {_i(g('mfe_session'))}, worst {C.pct(g('mae'))} on day "
               f"{_i(g('mae_session'))}; sold {C.day_label(g('exit_day'))} "
               f"{'close' if _s(g('exit_kind')) == 'close' else 'open (after a lower-circuit lock)'} at "
               f"{_px(g('exit_price'))}: net {C.pct(_div(g('net_bp'), 1e4), 2)} after costs - {res.upper()}")
        if C.fin(g("excess_bp")):
            s3 += f" ({C.pct(_div(g('excess_bp'), 1e4), 2)} vs buying everything that day)"
        s3 += "."
        s4 = []
        if _s(g("pre_home_index_name")) and C.fin(g("post_hold_home_ret")):
            s4.append(f"{_s(g('pre_home_index_name'))} {C.pct(g('post_hold_home_ret'))}")
        elif C.fin(g("post_hold_nifty_ret")):
            s4.append(f"NIFTY {C.pct(g('post_hold_nifty_ret'))}")
        if C.fin(g("post_hold_ad_ratio_avg")):
            s4.append(f"breadth averaged {float(g('post_hold_ad_ratio_avg')):.2f}")
        if C.fin(g("post_hold_stress_days")) and float(g("post_hold_stress_days")) > 0:
            s4.append(f"{_i(g('post_hold_stress_days'))} high-stress day(s)")
        s4t = ("During the hold " + ", ".join(s4)) if s4 else ""
        if _s(g("post_driver")) and C.fin(g("post_mkt_part_bp")) and C.fin(g("post_stock_part_bp")):
            s4t += ("; " if s4t else "") + (f"market part {C.pct(_div(g('post_mkt_part_bp'), 1e4), 2)}, stock part "
                                            f"{C.pct(_div(g('post_stock_part_bp'), 1e4), 2)}: "
                                            f"{C.DRIVER_TEXT.get(_s(g('post_driver')), _s(g('post_driver')))}")
        s3 += (" " + s4t + ".") if s4t else ""
    elif res == "Stuck":
        s3 = "Locked at the lower circuit for more than 20 sessions: no sale price (stuck)."
    else:
        s3 = f"No result yet ({_s(g('status')) or 'unknown'})."
    return " ".join([s1, s2, s3])


# ----------------------------------------------------------------------------------------------------------------
# column specs
# ----------------------------------------------------------------------------------------------------------------
class Col:
    """One journal column: header, width, number format, value getter, optional (fill, font, bold) styler, optional
    conditional format ('sign' = green / red text, or a SCALES key), and its meaning + source for the Read me."""

    def __init__(self, head: str, width: float, fmt: str, get: Callable[[dict], object], meaning: str, src: str,
                 style: Optional[Callable[[dict], Optional[Tuple]]] = None, cf: Optional[str] = None,
                 bold: bool = False, wrap: bool = False):
        self.head, self.width, self.fmt, self.get, self.meaning, self.src = head, width, fmt, get, meaning, src
        self.style, self.cf, self.bold, self.wrap = style, cf, bold, wrap


class Band:
    def __init__(self, name: str, colour: str, cols: List[Col], fold: bool = True):
        self.name, self.colour, self.cols, self.fold = name, colour, cols, fold


def _env_style(code_key: str):
    return lambda r: ENV_STYLE.get(_s(r.get(code_key)))


def trade_bands(include_all: bool, show_size: bool, other_engine: Dict[str, str]) -> List[Band]:
    g = lambda k: (lambda r: r.get(k))                                    # noqa: E731
    trade = [Col("Pick day", 15, "date", lambda r: _day(r.get("signal_day")),
                 "the session the engine picked the stock; the decision uses its close", "signal_day"),
             Col("Engine", 7, "text", lambda r: _s(r.get("engine")), "D or ENS", "engine"),
             Col("Rank", 5, "int", lambda r: _i(r.get("rank")), "the engine's rank that day (1-3 are traded)", "rank"),
             Col("Symbol", 13, "text", lambda r: _s(r.get("symbol")), "the stock", "symbol", bold=True)]
    if include_all:
        trade.append(Col("Traded", 7, "text", lambda r: "yes" if _b(r.get("traded")) else "",
                         "yes = rank 1-3, the paper plan's trade", "traded"))
    result = [
        Col("Result", 12, "text", result_of,
            "Winner = net above 0 after costs; Loser = 0 to -5%; Severe loser = -5% or worse; Stuck = locked at the "
            "lower circuit for more than 20 sessions", "net_bp, status", style=lambda r: RESULT_STYLE.get(result_of(r))),
        Col("Net %", 8, "pct2", lambda r: _div(r.get("net_bp"), 1e4), "the trade after the 0.35% round-trip cost",
            "net_bp / 100", cf="sign"),
        Col("vs buy-all %", 9, "pct2", lambda r: _div(r.get("excess_bp"), 1e4),
            "net minus buying every panel stock that day with the same plan (above 0 = beat the market)",
            "excess_bp / 100", cf="sign"),
        Col("Best %", 7, "pct1", lambda r: _f(r.get("mfe")), "the best high of days 1-5 against the buy price", "mfe",
            cf="best"),
        Col("Best day", 6, "int", lambda r: _i(r.get("mfe_session")), "the day (1-5) of that high", "mfe_session"),
        Col("Worst %", 7, "pct1", lambda r: _f(r.get("mae")), "the worst low of days 1-5 against the buy price", "mae",
            cf="worst"),
        Col("Worst day", 6, "int", lambda r: _i(r.get("mae_session")), "the day (1-5) of that low", "mae_session"),
        Col("+5% hit", 6, "int", lambda r: _i(r.get("touch_up5_session")),
            "the first day the high reached +5% (blank = never)", "touch_up5_session"),
        Col("-5% hit", 6, "int", lambda r: _i(r.get("touch_dn5_session")),
            "the first day the low reached -5% (blank = never)", "touch_dn5_session")]
    path = [Col("Gap %", 7, "pct1", lambda r: _f(r.get("entry_gap")),
                "the buy price (day-1 open) against the pick day's close", "entry_gap", cf="path")]
    for k in range(1, 6):
        path.append(Col(f"Day {k}", 7, "pct1", (lambda kk: (lambda r: _f(r.get(f"close_ret_s{kk}"))))(k),
                        f"day {k}'s close against the buy price", f"close_ret_s{k}", cf="path"))
    details = [Col("Bought", 15, "date", lambda r: _day(r.get("entry_day")), "the day bought, at its open", "entry_day"),
               Col("Buy", 9, "price", lambda r: _f(r.get("entry")), "the buy price (day-1 open)", "entry"),
               Col("Sold", 15, "date", lambda r: _day(r.get("exit_day")), "the day sold", "exit_day"),
               Col("Sell", 9, "price", lambda r: _f(r.get("exit_price")), "the sale price", "exit_price"),
               Col("Exit", 20, "text", exit_text,
                   "day-5 close (the plan), or the first free open after a lower-circuit lock", "exit_kind, sessions_held"),
               Col("Notes", 28, "text", lambda r: notes_text(r, other_engine),
                   "also traded by the other engine (its rank); repeat of = the same engine's pick of this stock still "
                   "inside its 5-day hold; not traded (with --all)",
                   "both_engines, other_rank, repeat_of, note")]
    stock = [
        Col("State", 19, "text", lambda r: state_label(r.get("pre_state")),
            "tl_state's primary state at the pick day's close", "pre_state",
            style=lambda r: STATE_STYLE.get(_s(r.get("pre_state"))) or (
                (None, "7F7F7F", False) if _s(r.get("pre_state")) else None)),
        Col("Confidence", 11.5, "text", lambda r: CONF_TEXT.get(_s(r.get("pre_state_confidence")), ""),
            "how well the evidence supports the state", "pre_state_confidence"),
        Col("Phase", 12, "text", lambda r: PHASE_TEXT.get(_s(r.get("pre_state_phase")), ""),
            "early (first 5 sessions) / mature / deteriorating", "pre_state_phase"),
        Col("Days in state", 7, "int", lambda r: _i(r.get("pre_state_age")),
            "sessions in the state, counted from the session the change began", "pre_state_age"),
        Col("Before", 17, "text", lambda r: state_label(r.get("pre_state_prev")), "the state before this one",
            "pre_state_prev"),
        Col("Pending", 15, "text", lambda r: state_label(r.get("pre_state_pending")),
            "a new reading not yet confirmed (tl_state confirms a change after 2 sessions)", "pre_state_pending"),
        Col("Event", 26, "text", lambda r: event_text(r.get("pre_state_event"), r.get("pre_state_event_age"), True),
            "breakout / breakdown / failed breakout in the last 10 sessions, its status and age", "pre_state_event"),
        Col("Flags", 34, "text", lambda r: flags_text(r.get("pre_state_risks"), r.get("pre_state_notes"), True),
            "risk flags and notes at the pick (overextended, divergences, large gaps ...)",
            "pre_state_risks, pre_state_notes"),
        Col("Volatility", 10, "text", lambda r: VOL_TEXT.get(_s(r.get("pre_state_vol_regime")), ""),
            "the stock's own volatility: compressed / normal / elevated / extreme", "pre_state_vol_regime",
            style=lambda r: VOL_STYLE.get(_s(r.get("pre_state_vol_regime")))),
        Col("Pick-day %", 8, "pct1", lambda r: _f(r.get("pre_stock_ret")), "the stock's own move on the pick day",
            "pre_stock_ret", cf="sign"),
        Col("20-day %", 8, "pct1", lambda r: _f(r.get("pre_state_r20")), "the stock's 20-session return",
            "pre_state_r20", cf="sign"),
        Col("RSI 14", 6, "num0", lambda r: _f(r.get("pre_state_rsi14")), "RSI 14 (Wilder)", "pre_state_rsi14"),
        Col("Moves with", 12, "text", lambda r: _s(r.get("pre_home_index_name")),
            "the index the stock follows most (its home index)", "pre_home_index_name"),
        Col("Beta", 6, "num2", lambda r: _f(r.get("pre_beta_home")), "beta to that index", "pre_beta_home")]
    if show_size:
        stock.append(Col("Size", 8, "text", lambda r: _s(r.get("pre_size_class")), "AMFI size class", "pre_size_class"))
    market = [
        Col("Environment", 15, "text", lambda r: C.ENV_TEXT.get(_s(r.get("pre_environment")), ""),
            "tradelab's market environment at the pick day's close (HIGH STRESS overrides the direction)",
            "pre_environment", style=_env_style("pre_environment")),
        Col("Stress reason", 26, "text", lambda r: _s(r.get("pre_stress_reason")), "why the day was high stress",
            "pre_stress_reason"),
        Col("NIFTY day %", 8, "pct1", g("pre_nifty_ret"), "NIFTY on the pick day", "pre_nifty_ret", cf="sign"),
        Col("Smallcap day %", 9.5, "pct1", g("pre_small_ret"), "Smallcap 250 on the pick day", "pre_small_ret", cf="sign"),
        Col("NIFTY 20d %", 8, "pct1", g("pre_nifty_ret20"), "NIFTY over the last 20 sessions", "pre_nifty_ret20",
            cf="sign"),
        Col("Breadth", 8.5, "ratio", g("pre_ad_ratio"),
            "advancers / (advancers + decliners) among the cached stocks on the pick day (0.5 = even)",
            "pre_ad_ratio", cf="breadth"),
        Col("Breadth 10d", 8.5, "ratio", g("pre_ad_ratio_10"), "the same, 10-session average", "pre_ad_ratio_10",
            cf="breadth"),
        Col("VIX", 6, "num1", g("pre_vix"), "India VIX at the pick day's close", "pre_vix"),
        Col("VIX 5d %", 7, "pct1", g("pre_vix_chg5_pct"), "VIX change over 5 sessions (rising = red)",
            "pre_vix_chg5_pct", cf="vix5"),
        Col("VIX pctile", 7, "pct0", g("pre_vix_pct250"), "where VIX sits in its last 250 sessions (high = red)",
            "pre_vix_pct250", cf="vixpct")]
    for c in market[2:]:
        if c.fmt in ("pct1", "ratio", "num1", "pct0"):
            c.get = (lambda f: (lambda r: _f(f(r))))(c.get)
    hold = [
        Col("What drove it", 34, "text", lambda r: C.DRIVER_TEXT.get(_s(r.get("post_driver")), _s(r.get("post_driver"))),
            "market (market part at least 2x the stock part) / stock / both / opposite ways / flat (all under 0.5%)",
            "post_driver"),
        Col("Market part %", 8, "pct2", lambda r: _div(r.get("post_mkt_part_bp"), 1e4),
            "beta x the home index over the hold: what the market gave", "post_mkt_part_bp / 100", cf="sign"),
        Col("Stock part %", 8, "pct2", lambda r: _div(r.get("post_stock_part_bp"), 1e4),
            "gross return minus the market part: what the stock did on its own", "post_stock_part_bp / 100", cf="sign"),
        Col("Home index %", 8, "pct1", lambda r: _f(r.get("post_hold_home_ret")), "the stock's home index over the hold",
            "post_hold_home_ret", cf="sign"),
        Col("NIFTY hold %", 8, "pct1", lambda r: _f(r.get("post_hold_nifty_ret")), "NIFTY over the hold",
            "post_hold_nifty_ret", cf="sign"),
        Col("Smallcap hold %", 9.5, "pct1", lambda r: _f(r.get("post_hold_small_ret")), "Smallcap 250 over the hold",
            "post_hold_small_ret", cf="sign"),
        Col("Avg breadth", 8.5, "ratio", lambda r: _f(r.get("post_hold_ad_ratio_avg")),
            "breadth averaged over the days held", "post_hold_ad_ratio_avg", cf="breadth"),
        Col("Stress days", 6, "int", lambda r: _i(r.get("post_hold_stress_days")), "days held that were high stress",
            "post_hold_stress_days"),
        Col("Env. entry day", 15, "text", lambda r: C.ENV_TEXT.get(_s(r.get("post_entry_environment")), ""),
            "the environment at the close of the day bought (hindsight)", "post_entry_environment",
            style=_env_style("post_entry_environment"))]
    stry = [Col("Story", 120, "text", story, "the pick in one line: the market that day, the stock and its state, the "
                                            "trade, the hold", "built from the columns above")]
    return [Band("TRADE", "D9D9D9", trade, fold=False), Band("RESULT", "BDD7EE", result, fold=False),
            Band("CLOSE vs BUY PRICE, DAY BY DAY", "DDEBF7", path, fold=False), Band("TRADE DETAILS", "E7E6E6", details),
            Band("STOCK AT THE PICK", "C6E0B4", stock), Band("MARKET AT THE PICK", "FFE699", market),
            Band("DURING THE HOLD (hindsight)", "F8CBAD", hold), Band("STORY", "D9D2E9", stry, fold=False)]


# ----------------------------------------------------------------------------------------------------------------
# the workbook writer (openpyxl write-only: fast, low memory)
# ----------------------------------------------------------------------------------------------------------------
class Styles:
    def __init__(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        self.WOC = WriteOnlyCell
        self.Font, self.PatternFill, self.Alignment, self.Border, self.Side = Font, PatternFill, Alignment, Border, Side
        self._fonts, self._fills = {}, {}
        self.left_edge = Border(left=Side(style="thin", color="A6A6A6"))
        self.no_border = Border()
        self.head_border = Border(bottom=Side(style="thin", color="7F7F7F"))
        self.head_edge = Border(bottom=Side(style="thin", color="7F7F7F"), left=Side(style="thin", color="A6A6A6"))
        self.wrap_c = Alignment(wrap_text=True, vertical="center", horizontal="center")
        self.left = Alignment(vertical="top")
        self.center = Alignment(horizontal="center", vertical="top")

    def font(self, colour: str = "000000", bold: bool = False, italic: bool = False, size: float = 10):
        k = (colour, bold, italic, size)
        if k not in self._fonts:
            self._fonts[k] = self.Font(name=FONT, size=size, color=colour, bold=bold, italic=italic)
        return self._fonts[k]

    def fill(self, colour: Optional[str]):
        if not colour:
            return None
        if colour not in self._fills:
            self._fills[colour] = self.PatternFill("solid", fgColor=colour)
        return self._fills[colour]


def _cell(ws, S: Styles, value, fmt: Optional[str] = None, fill: Optional[str] = None, font=None, border=None,
          align=None):
    c = S.WOC(ws, value=value)
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"                        # text that looks like a formula stays text
    c.font = font or S.font()
    if fmt and NF.get(fmt):
        c.number_format = NF[fmt]
    f = S.fill(fill)
    if f is not None:
        c.fill = f
    if border is not None:
        c.border = border
    if align is not None:
        c.alignment = align
    return c


def write_table(wb, S: Styles, name: str, title: str, bands: List[Band], rows: List[dict], freeze_after: str,
                zebra_key: Optional[Callable[[dict], object]] = None, active: bool = False) -> dict:
    """Row 1 title, row 2 the bands, row 3 the headers (with filters), data from row 4. Returns what it wrote."""
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(name)
    cols: List[Tuple[Band, Col, bool]] = []
    for b in bands:
        for j, c in enumerate(b.cols):
            cols.append((b, c, j == 0))
    ncol = len(cols)
    last = get_column_letter(ncol)
    for k, (b, c, first) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(k)].width = c.width
        if b.fold and len(b.cols) > 1 and not first:
            # a foldable band keeps its first (key) column; the rest fold. Adjacent columns at one outline level are
            # ONE group in Excel, so the key column (level 0) is what separates the bands' groups.
            ws.column_dimensions[get_column_letter(k)].outline_level = 1
    if any(b.fold and len(b.cols) > 1 for b in bands):
        from openpyxl.worksheet.properties import Outline
        ws.sheet_properties.outlinePr = Outline(summaryBelow=True, summaryRight=False)   # the [-] sits over the key
    heads = [c.head for _, c, _ in cols]
    fz = heads.index(freeze_after) + 2 if freeze_after in heads else 2
    ws.freeze_panes = f"{get_column_letter(fz)}4"
    ws.append([_cell(ws, S, title, font=S.font(bold=True, size=11))])
    ws.append([_cell(ws, S, b.name if first else None, fill=b.colour, font=S.font(bold=True),
                     border=S.left_edge if first else None) for b, c, first in cols])
    ws.append([_cell(ws, S, c.head, fill=b.colour, font=S.font(bold=True), align=S.wrap_c,
                     border=S.head_edge if first else S.head_border) for b, c, first in cols])
    zeb, prev, n = False, object(), 0
    for r in rows:
        if zebra_key is not None:
            key = zebra_key(r)
            if key != prev:
                zeb, prev = (not zeb) if n else False, key
        out = []
        for b, c, first in cols:
            v = c.get(r)
            st = c.style(r) if c.style else None
            fill = st[0] if st and st[0] else (ZEBRA if zeb else None)
            font = S.font(st[1], bold=st[2] or c.bold, italic=(st[0] is None)) if st else S.font(bold=c.bold)
            out.append(_cell(ws, S, v, c.fmt, fill, font, S.left_edge if first else None,
                             S.center if c.fmt in ("int", "text") and c.width <= 7 else S.left))
        ws.append(out)
        n += 1
    if n:
        ws.auto_filter.ref = f"A3:{last}{3 + n}"
        for k, (b, c, _) in enumerate(cols, start=1):
            if not c.cf:
                continue
            L = get_column_letter(k)
            rng = f"{L}4:{L}{3 + n}"
            if c.cf == "sign":
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                              font=S.Font(name=FONT, size=10, color=GOOD_TXT)))
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                              font=S.Font(name=FONT, size=10, color=BAD_TXT)))
            else:
                pts = SCALES[c.cf]
                if len(pts) == 2:
                    rule = ColorScaleRule(start_type="num", start_value=pts[0][0], start_color=pts[0][1],
                                          end_type="num", end_value=pts[1][0], end_color=pts[1][1])
                else:
                    rule = ColorScaleRule(start_type="num", start_value=pts[0][0], start_color=pts[0][1],
                                          mid_type="num", mid_value=pts[1][0], mid_color=pts[1][1],
                                          end_type="num", end_value=pts[2][0], end_color=pts[2][1])
                ws.conditional_formatting.add(rng, rule)
    return {"sheet": name, "rows": n, "cols": ncol, "freeze": ws.freeze_panes, "columns": heads, "ws": ws}


# ----------------------------------------------------------------------------------------------------------------
# the export
# ----------------------------------------------------------------------------------------------------------------
def find_export(spec: str) -> Path:
    if not spec:
        raise TLError("--export is required: an export folder (TLX_YYYYMMDD_NNN, or its full path)")
    p = Path(spec)
    if not (p / "trades.parquet").exists() and not p.is_absolute():
        base = os.environ.get("TRADELAB_OUT")
        if base and (Path(base) / "export" / spec / "trades.parquet").exists():
            p = Path(base) / "export" / spec
    if not (p / "trades.parquet").exists():
        hint = "" if os.environ.get("TRADELAB_OUT") else " (start with env_tradelab.bat so TRADELAB_OUT is set)"
        raise TLError(f"{spec} is not an export folder: no trades.parquet there{hint}")
    return p


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def check_trades_hash(ex: Path) -> str:
    """trades.parquet must match the export's own outputs.sha256 (a half-copied or edited export is refused)."""
    h = sha256_file(ex / "trades.parquet")
    f = ex / "outputs.sha256"
    if not f.exists():
        return f"{h[:16]} (the export has no outputs.sha256 to check it against)"
    listed = {}
    for line in f.read_text(encoding="utf-8", errors="replace").splitlines():
        parts = line.strip().split(None, 1)
        if len(parts) == 2:
            listed[parts[1].strip().replace("\\", "/")] = parts[0].lower()
    if "trades.parquet" not in listed:
        return f"{h[:16]} (outputs.sha256 does not list trades.parquet)"
    if listed["trades.parquet"] != h:
        raise TLError(f"{ex / 'trades.parquet'} does not match the export's outputs.sha256 (file {h[:16]}, listed "
                      f"{listed['trades.parquet'][:16]}): the export was changed or copied half-way. Copy it again "
                      f"or re-run the export.")
    return f"{h[:16]} (matches the export's outputs.sha256)"


USED = ("trade_id", "engine", "signal_day", "symbol", "rank", "traded", "both_engines", "other_rank", "repeat",
        "repeat_of", "status", "note", "entry_day", "entry", "entry_gap", "exit_day", "exit_price", "exit_kind",
        "sessions_held", "net_bp", "excess_bp", "mkt_net_bp", "mfe", "mae", "mfe_session", "mae_session",
        "touch_up5_session", "touch_dn5_session", "close_ret_s1", "close_ret_s2", "close_ret_s3", "close_ret_s4",
        "close_ret_s5", "pre_state", "pre_state_confidence", "pre_state_phase", "pre_state_age", "pre_state_prev",
        "pre_state_pending", "pre_state_event", "pre_state_event_age", "pre_state_risks", "pre_state_notes",
        "pre_state_vol_regime", "pre_stock_ret", "pre_state_r20", "pre_state_rsi14", "pre_home_index_name",
        "pre_beta_home", "pre_size_class", "pre_environment", "pre_stress_reason", "pre_nifty_ret", "pre_small_ret",
        "pre_nifty_ret20", "pre_ad_ratio", "pre_ad_ratio_10", "pre_vix", "pre_vix_chg5_pct", "pre_vix_pct250",
        "post_entry_environment", "post_hold_home_ret", "post_hold_nifty_ret", "post_hold_small_ret",
        "post_hold_ad_ratio_avg", "post_hold_stress_days", "post_mkt_part_bp", "post_stock_part_bp", "post_driver")
NEEDED = ("engine", "signal_day", "symbol", "rank", "net_bp", "status")


def load_trades(ex: Path) -> Tuple[pd.DataFrame, List[str]]:
    import pyarrow.parquet as pq
    have = set(pq.ParquetFile(ex / "trades.parquet").schema_arrow.names)
    miss_need = [c for c in NEEDED if c not in have]
    if miss_need:
        raise TLError(f"{ex / 'trades.parquet'} has no {', '.join(miss_need)}: not a tradelab export")
    T = pd.read_parquet(ex / "trades.parquet", columns=[c for c in USED if c in have])
    missing = [c for c in USED if c not in have]
    for c in missing:
        T[c] = np.nan
    if "traded" in missing:
        T["traded"] = T["rank"].astype(float) <= 3
    return T, missing


def order_rows(T: pd.DataFrame) -> pd.DataFrame:
    eng = {e: i for i, e in enumerate(ENGINE_ORDER)}
    X = T.copy()
    X["_d"] = pd.to_datetime(X["signal_day"])
    X["_e"] = X["engine"].map(lambda e: eng.get(e, len(eng)))
    X = X.sort_values(["_d", "_e", "engine", "rank"], ascending=[False, True, True, True], kind="mergesort")
    return X.drop(columns=["_d", "_e"]).reset_index(drop=True)


def market_rows(ex: Path, T: pd.DataFrame, engines: Sequence[str]) -> Tuple[Optional[List[dict]], str]:
    p = ex / "market_days.parquet"
    if not p.exists():
        return None, "market_days.parquet is not in this export (exported with --context none)"
    M = pd.read_parquet(p)
    M["date"] = pd.to_datetime(M["date"])
    X = T[T["traded"].map(_b)].copy()
    X["_d"] = pd.to_datetime(X["signal_day"])
    X["_net"] = X["net_bp"].astype(float) / 1e4
    for e in engines:
        G = X[X["engine"] == e].groupby("_d")
        M[f"_net_{e}"] = M["date"].map(G["_net"].mean())
        M[f"_win_{e}"] = M["date"].map(G["_net"].apply(lambda v: f"{int((v > 0).sum())} of {int(v.notna().sum())}"))
    M["_mkt"] = M["date"].map(X.groupby("_d")["mkt_net_bp"].first().astype(float) / 1e4)
    M = M.sort_values("date", ascending=False).reset_index(drop=True)
    return M.to_dict("records"), ""


def market_bands(engines: Sequence[str]) -> List[Band]:
    f = lambda k: (lambda r: _f(r.get(k)))                                # noqa: E731
    day = [Col("Date", 15, "date", lambda r: _day(r.get("date")), "the session", "date")]
    mk = [Col("Environment", 15, "text", lambda r: C.ENV_TEXT.get(_s(r.get("environment")), ""),
              "tradelab's market environment at the close", "environment", style=_env_style("environment")),
          Col("Stress reason", 26, "text", lambda r: _s(r.get("stress_reason")), "why the day was high stress",
              "stress_reason"),
          Col("NIFTY %", 7, "pct1", f("nifty_ret"), "NIFTY that day", "nifty_ret", cf="sign"),
          Col("Midcap %", 7, "pct1", f("mid_ret"), "Midcap 150 that day", "mid_ret", cf="sign"),
          Col("Smallcap %", 9, "pct1", f("small_ret"), "Smallcap 250 that day", "small_ret", cf="sign"),
          Col("NIFTY 20d %", 8, "pct1", f("nifty_ret20"), "NIFTY over 20 sessions", "nifty_ret20", cf="sign"),
          Col("Smallcap 20d %", 9, "pct1", f("small_ret20"), "Smallcap 250 over 20 sessions", "small_ret20", cf="sign"),
          Col("NIFTY vs 50d avg", 9, "pct1", f("nifty_vs_ma50"), "NIFTY against its 50-session average",
              "nifty_vs_ma50", cf="sign")]
    br = [Col("Breadth", 8.5, "ratio", f("ad_ratio"), "advancers / (advancers + decliners), cached stocks", "ad_ratio",
              cf="breadth"),
          Col("Breadth 10d", 8.5, "ratio", f("ad_ratio_10"), "10-session average", "ad_ratio_10", cf="breadth"),
          Col("Up", 6, "count", f("adv"), "stocks up", "adv"), Col("Down", 6, "count", f("dec"), "stocks down", "dec"),
          Col("VIX", 6, "num1", f("vix"), "India VIX", "vix"),
          Col("VIX day %", 7, "pct1", f("vix_chg_pct"), "VIX change that day", "vix_chg_pct", cf="vix5"),
          Col("VIX 5d %", 7, "pct1", f("vix_chg5_pct"), "VIX change over 5 sessions", "vix_chg5_pct", cf="vix5"),
          Col("VIX pctile", 7, "pct0", f("vix_pct250"), "VIX percentile in 250 sessions", "vix_pct250", cf="vixpct")]
    pk = []
    for e in engines:
        pk.append(Col(f"{e} top-3 net %", 9, "pct2", (lambda ee: (lambda r: _f(r.get(f"_net_{ee}"))))(e),
                      f"the average net of {e}'s three picks made that day (bought the next morning)",
                      "trades: net_bp", cf="sign"))
        pk.append(Col(f"{e} winners", 8.5, "text", (lambda ee: (lambda r: _s(r.get(f"_win_{ee}"))))(e),
                      f"how many of {e}'s picks that day ended above 0 after costs", "trades: net_bp"))
    pk.append(Col("Buy-all net %", 9, "pct2", lambda r: _f(r.get("_mkt")),
                  "buying every panel stock that day with the same plan", "trades: mkt_net_bp", cf="sign"))
    sm = [Col("Summary", 120, "text", lambda r: _s(r.get("summary")), "tradelab's one-line summary of the day",
              "summary")]
    return [Band("DAY", "D9D9D9", day, fold=False), Band("MARKET", "FFE699", mk), Band("BREADTH AND VIX", "FCE4D6", br),
            Band("PICKS MADE THIS DAY", "BDD7EE", pk), Band("SUMMARY", "D9D2E9", sm, fold=False)]


DIM_ORDER = ("environment at the pick", "high stress at the pick", "direction at the pick (stress ignored)",
             "breadth at the pick (10-day A/D)", "VIX 5-day change at the pick", "NIFTY minus Smallcap, 20 days, at the pick",
             "stock state at the pick", "state phase at the pick", "state confidence at the pick",
             "event in the last 10 sessions", "S05 overextension at the pick", "stock volatility at the pick",
             "home index", "size class (AMFI)", "environment on the entry day", "Smallcap over the hold",
             "what drove the trade")


def bucket_text(dim: str, b: str) -> str:
    b = _s(b)
    if "environment" in dim or dim.startswith("direction"):
        return C.ENV_TEXT.get(b, b)
    if dim == "stock state at the pick":
        return state_label(b) or b
    if dim == "event in the last 10 sessions":
        return event_text(b, None, True) or b
    if dim == "what drove the trade":
        return C.DRIVER_TEXT.get(b, b)
    if dim in ("state phase at the pick",):
        return PHASE_TEXT.get(b, b)
    if dim in ("state confidence at the pick",):
        return CONF_TEXT.get(b, b.capitalize() if b.isupper() else b)
    if dim == "stock volatility at the pick":
        return VOL_TEXT.get(b, b)
    return b


def verdict(lo, hi) -> str:
    if not (C.fin(lo) and C.fin(hi)):
        return ""
    return "above 0" if float(lo) > 0 else ("below 0" if float(hi) < 0 else "spans 0")


def worked_rows(ex: Path, engines: Sequence[str]) -> Tuple[Optional[List[dict]], str]:
    p = ex / "context_breakdown.csv"
    if not p.exists():
        return None, "context_breakdown.csv is not in this export (exported with --context none)"
    B = pd.read_csv(p)
    eng = {e: i for i, e in enumerate(list(engines) + [e for e in ENGINE_ORDER if e not in engines])}
    dims = {d: i for i, d in enumerate(DIM_ORDER)}
    B["_e"] = B["engine"].map(lambda e: eng.get(e, 99))
    B["_t"] = B["timing"].map(lambda t: 0 if t == "before" else 1)
    B["_dm"] = B["dimension"].map(lambda d: dims.get(d, 50))
    B["_k"] = np.arange(len(B))
    B = B.sort_values(["_e", "_t", "_dm", "dimension", "_k"], kind="mergesort").reset_index(drop=True)
    return B.to_dict("records"), ""


def worked_bands() -> List[Band]:
    f = lambda k: (lambda r: _f(r.get(k)))                                # noqa: E731
    grp = [Col("Engine", 7, "text", lambda r: _s(r.get("engine")), "D or ENS (traded picks, ranks 1-3)", "engine"),
           Col("Known", 12, "text", lambda r: "at the pick" if r.get("timing") == "before" else "afterwards",
               "at the pick = usable for a decision; afterwards = hindsight, cannot be traded on", "timing"),
           Col("Grouping", 40, "text", lambda r: _s(r.get("dimension")), "what the trades are split by", "dimension"),
           Col("Group", 24, "text", lambda r: bucket_text(_s(r.get("dimension")), r.get("bucket")), "the bucket",
               "bucket", bold=True)]
    size = [Col("Trades", 7, "count", f("trades"), "trades in the bucket", "trades"),
            Col("Days", 6, "count", f("days"), "pick days with a trade in the bucket", "days"),
            Col("Share", 7, "share", f("share"), "share of the engine's trades", "share")]
    rates = [Col("Win rate", 7, "pct0", f("win_rate"), "share with net above 0", "win_rate", cf="rate"),
             Col("Severe-loser rate", 8, "pct0", f("severe_rate"), "share losing 5% or more", "severe_rate", cf="rate")]
    net = [Col("Net bp", 7, "bp", f("net_bp_mean"), "mean net per trade, bp (100 bp = 1%)", "net_bp_mean", cf="sign"),
           Col("low", 7, "bp", f("net_bp_lo"), "95% interval, low end", "net_bp_lo"),
           Col("high", 7, "bp", f("net_bp_hi"), "95% interval, high end", "net_bp_hi"),
           Col("Net vs 0", 9, "text", lambda r: verdict(r.get("net_bp_lo"), r.get("net_bp_hi")),
               "the interval above 0, below 0, or spanning 0", "net_bp_lo, net_bp_hi",
               style=lambda r: VERDICT_STYLE.get(verdict(r.get("net_bp_lo"), r.get("net_bp_hi"))))]
    exc = [Col("vs buy-all bp", 8, "bp", f("excess_bp_mean"), "mean net minus buying everything that day, bp",
               "excess_bp_mean", cf="sign"),
           Col("low", 7, "bp", f("excess_bp_lo"), "95% interval, low end", "excess_bp_lo"),
           Col("high", 7, "bp", f("excess_bp_hi"), "95% interval, high end", "excess_bp_hi"),
           Col("Beat buy-all?", 9, "text", lambda r: verdict(r.get("excess_bp_lo"), r.get("excess_bp_hi")),
               "the interval above 0 (beat it), below 0, or spanning 0 (no clear difference)",
               "excess_bp_lo, excess_bp_hi",
               style=lambda r: VERDICT_STYLE.get(verdict(r.get("excess_bp_lo"), r.get("excess_bp_hi"))))]
    parts = [Col("Market part bp", 8, "bp", f("post_mkt_part_bp_mean"), "mean market part (beta x home index)",
                 "post_mkt_part_bp_mean", cf="sign"),
             Col("Stock part bp", 8, "bp", f("post_stock_part_bp_mean"), "mean stock part (gross minus market part)",
                 "post_stock_part_bp_mean", cf="sign")]
    return [Band("GROUP", "D9D9D9", grp, fold=False), Band("SIZE", "E7E6E6", size),
            Band("RATES", "DDEBF7", rates), Band("NET PER TRADE", "BDD7EE", net, fold=False),
            Band("AGAINST BUYING EVERYTHING", "C6E0B4", exc, fold=False), Band("MARKET vs STOCK", "FCE4D6", parts)]


# ----------------------------------------------------------------------------------------------------------------
# the Read me sheet
# ----------------------------------------------------------------------------------------------------------------
def _hex_mix(a: str, b: str, t: float) -> str:
    pa = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    pb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(pa, pb))


def _scale_colour(points, v: float) -> str:
    if v <= points[0][0]:
        return points[0][1]
    for (x0, c0), (x1, c1) in zip(points[:-1], points[1:]):
        if v <= x1:
            return _hex_mix(c0, c1, (v - x0) / (x1 - x0) if x1 > x0 else 0.0)
    return points[-1][1]


def write_readme(wb, S: Styles, info: dict, sheets: List[Tuple[str, List[Band]]]) -> None:
    ws = wb.create_sheet("Read me")
    for k, w in enumerate((22, 26, 22, 70, 34), start=1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(k)].width = w
    B = lambda t, size=10: _cell(ws, S, t, font=S.font(bold=True, size=size))          # noqa: E731
    T = lambda t: _cell(ws, S, t)                                                       # noqa: E731
    ws.append([B(f"Trade journal - export {info['export']}", 14)])
    ws.append([T(f"Made {info['made']} by {TOOL} ({VERSION}) from {info['path']}.")])
    ws.append([T(f"trades.parquet {info['hash']}.")])
    ws.append([])
    ws.append([B("THE PLAN", 11)])
    for line in ("Buy at the open of the session after the pick day, sell at the close of the 5th session. Costs: 0.35% "
                 "round trip.",
                 "A close locked at the lower circuit sells at the first free open, up to 20 sessions later (column Exit).",
                 f"Result: Winner = net above 0 after costs; Loser = 0 to -5%; Severe loser = -5% or worse (tradelab's "
                 f"{C.SEVERE_BP:+.0f} bp rule); Stuck = locked for more than 20 sessions, no sale price."):
        ws.append([T(line)])
    ws.append([])
    ws.append([B("THE SHEETS", 11)])
    for nm, txt in (("Trades", info["trades_line"]),
                    ("Market days", info["market_line"]),
                    ("What worked", info["worked_line"]),
                    ("Read me", "this page: the colours, then every column in plain words")):
        ws.append([_cell(ws, S, nm, font=S.font(bold=True)), T(txt)])
    ws.append([])
    ws.append([B("HOW TO READ THEM", 11)])
    for line in ("Row 3 of each sheet has filter arrows: filter by engine, result, state, environment, any column.",
                 "The left columns stay in place while you scroll right (Trades: up to Net %).",
                 "The [-] buttons above the bands fold a band down to its key column (State, Environment, What drove "
                 "it ...); [+] opens it again.",
                 "Trades: the picks of one day sit together; days alternate white and light grey. Market days: months "
                 "alternate.",
                 "Numbers: green text = above 0, red text = below 0. Percent columns are true percentages (0.35% = 35 "
                 "bp).",
                 "The Story column (last) tells each pick in one line; widen it or read it in the formula bar."):
        ws.append([T(line)])
    ws.append([])
    ws.append([B("COLOURS", 11)])

    def legend(label: str, items: List[Tuple[str, Tuple, str]]):
        """One row per sample: the label (first row only), the coloured sample, what it means."""
        for j, (txt, st, meaning) in enumerate(items):
            fill, fc, bold = st
            ws.append([_cell(ws, S, label if j == 0 else None, font=S.font(bold=True)),
                       _cell(ws, S, txt, fill=fill, font=S.font(fc, bold=bold, italic=fill is None and fc == "7F7F7F"),
                             align=S.center), T(meaning)])
        ws.append([])
    plain = (None, "000000", False)
    legend("Result", [("Winner", RESULT_STYLE["Winner"], "net above 0 after costs"),
                      ("Loser", RESULT_STYLE["Loser"], "net 0 to -5% after costs"),
                      ("Severe loser", RESULT_STYLE["Severe loser"],
                       f"net -5% or worse after costs (tradelab's {C.SEVERE_BP:+.0f} bp rule)"),
                      ("Stuck", RESULT_STYLE["Stuck"], "locked at the lower circuit for more than 20 sessions: no sale")])
    legend("Environment", [(C.ENV_TEXT[k], ENV_STYLE[k], C.DEFINITIONS.get(k, "")) for k in ENV_STYLE] +
           [("", plain, f"'up' = the index closes at or above its {C.MA_DAYS}-session average; breadth strong = "
                        f"{C.AD_AVG_DAYS}-day advancers / (advancers + decliners) at or above {C.AD_STRONG:.2f}")])
    legend("Stock state", [(state_label(k), STATE_STYLE[k], ST.STATE_NAMES.get(k, k)) for k in STATE_STYLE] +
           [(state_label("UNKNOWN"), (None, "7F7F7F", False),
             "grey italic: no state is clear (evidence conflicts), or too little data")])
    legend("Volatility", [(VOL_TEXT[k], VOL_STYLE[k], f"the stock's own volatility is {VOL_TEXT[k].lower()}")
                          for k in VOL_STYLE] + [("Normal", plain, "no colour")])
    legend("Day 1-5, Gap", [(f"{v:+.0%}", (_scale_colour(SCALES["path"], v), "000000", False),
                             "close against the buy price: red below, green above, full colour at -8% / +8%"
                             if v == -0.08 else "") for v in (-0.08, -0.04, 0.0, 0.04, 0.08)])
    legend("Breadth", [(f"{v:.2f}", (_scale_colour(SCALES["breadth"], v), "000000", False),
                        "share of stocks up: 0.50 = as many up as down; full colour at 0.25 / 0.75" if v == 0.25 else "")
                       for v in (0.25, 0.4, 0.5, 0.6, 0.75)])
    legend("VIX", [(f"5-day {v:+.0%}", (_scale_colour(SCALES["vix5"], v), "000000", False),
                    "rising VIX (fear) red, falling green; VIX percentile turns red above the 50th" if v == -0.3 else "")
                   for v in (-0.3, 0.0, 0.3)])
    legend("Numbers", [("+1.25%", (None, GOOD_TXT, False), "green text: above 0"),
                       ("-1.25%", (None, BAD_TXT, False), "red text: below 0")])
    legend("What worked", [(k, v, m) for (k, v), m in zip(VERDICT_STYLE.items(), (
        "the 95% interval is above 0", "the 95% interval is below 0", "the interval spans 0: no clear difference"))])
    ws.append([])
    ws.append([B("EVERY COLUMN", 11)])
    ws.append([_cell(ws, S, h, font=S.font(bold=True), fill="D9D9D9")
               for h in ("Sheet", "Band", "Column", "Meaning", "From the export")])
    for sname, bands in sheets:
        for b in bands:
            for c in b.cols:
                ws.append([T(sname), _cell(ws, S, b.name, fill=b.colour), _cell(ws, S, c.head, font=S.font(bold=True)),
                           T(c.meaning), T(c.src)])
    ws.append([])
    ws.append([B("NOTES", 11)])
    for line in info["notes"]:
        ws.append([T(line)])


# ----------------------------------------------------------------------------------------------------------------
# the run
# ----------------------------------------------------------------------------------------------------------------
def build(ex: Path, include_all: bool = False, name: str = DEFAULT_NAME, log: Callable[[str], None] = say) -> dict:
    _need_openpyxl()
    from openpyxl import Workbook
    t0 = time.time()
    ex = Path(ex)
    hash_line = check_trades_hash(ex)
    T, missing = load_trades(ex)
    engines = [e for e in ENGINE_ORDER if e in set(T["engine"])] + sorted(set(T["engine"]) - set(ENGINE_ORDER))
    other = {}
    if len(engines) == 2:
        other = {engines[0]: engines[1], engines[1]: engines[0]}
    sizes = set(T["pre_size_class"].dropna().astype(str))
    show_size = bool(sizes & {"large", "mid", "small"})
    X = T if include_all else T[T["traded"].map(_b)]
    X = order_rows(X)
    rows = X.to_dict("records")
    days = X["signal_day"].nunique()
    d0, d1 = (pd.Timestamp(X["signal_day"].min()).date(), pd.Timestamp(X["signal_day"].max()).date()) if len(X) else ("", "")
    what = "picks (ranks 1-10, traded = 1-3)" if include_all else "traded picks (ranks 1-3)"
    tb = trade_bands(include_all, show_size, other)
    log(f"{TOOL} ({VERSION}) | export {ex} | {len(X):,} {what} of {', '.join(engines)} on {days:,} days")
    S = Styles()
    wb = Workbook(write_only=True)
    sheet_name = "Picks" if include_all else "Trades"
    title = (f"{'Every pick' if include_all else 'Traded picks'}, newest first - {len(X):,} on {days:,} pick days, {d0} "
             f"to {d1} - export {ex.name}. Colours and columns: see the Read me sheet.")
    res = {"trades": write_table(wb, S, sheet_name, title, tb, rows, "Net %", zebra_key=lambda r: r.get("signal_day"))}
    sheets = [(sheet_name, tb)]
    M, mnote = market_rows(ex, T, engines)
    if M is not None:
        mb = market_bands(engines)
        mt = (f"Every session, newest first - {len(M):,} sessions. 'Picks made this day' = the picks chosen at that "
              f"close (bought the next morning).")
        res["market"] = write_table(wb, S, "Market days", mt, mb, M, "Date",
                                    zebra_key=lambda r: (pd.Timestamp(r["date"]).year, pd.Timestamp(r["date"]).month))
        sheets.append(("Market days", mb))
    else:
        ws = wb.create_sheet("Market days")
        ws.append([_cell(ws, S, mnote, font=S.font(bold=True))])
    W, wnote = worked_rows(ex, engines)
    if W is not None:
        wbands = worked_bands()
        wt = ("What worked - tradelab's own breakdown of the traded picks (context_breakdown.csv), copied as it is. 95% "
              "block-bootstrap intervals over trading days. Descriptive only: many groups are compared at once, so one "
              "interval clear of 0 is not proof; 'afterwards' rows are hindsight.")
        res["worked"] = write_table(wb, S, "What worked", wt, wbands, W, "Group",
                                    zebra_key=lambda r: (r.get("engine"), r.get("dimension")))
        sheets.append(("What worked", wbands))
    else:
        ws = wb.create_sheet("What worked")
        ws.append([_cell(ws, S, wnote, font=S.font(bold=True))])
    notes = ["Values, not formulas: this file is a view of the export. Delete it and run tl_journal.py again any time; "
             "the export's other files are not touched (outputs.sha256 does not list this file).",
             "Breadth counts the stocks in the cache (survivors), not NSE's official advance / decline.",
             "'During the hold' columns are hindsight: they explain a trade, they could not have chosen it."]
    if not show_size:
        notes.append("Size class is not shown: this export has no AMFI size lists (every pick reads "
                     f"'{sorted(sizes)[0] if sizes else 'blank'}').")
    if missing:
        notes.append(f"This export has no {', '.join(missing)}: those columns are blank.")
    for nt in (mnote, wnote):
        if nt:
            notes.append(nt)
    info = {"export": ex.name, "path": str(ex), "hash": hash_line, "made": time.strftime("%Y-%m-%d %H:%M"),
            "notes": notes,
            "trades_line": (f"{len(X):,} {what}, newest first; inside a day {' before '.join(engines)}, rank 1 first. "
                            f"The trade, its result, the close day by day, the stock's state and the market at the pick, "
                            f"the hold, and the story"),
            "market_line": (f"{len(M):,} sessions, newest first: environment, indices, breadth, VIX, and how that day's "
                            f"picks did" if M is not None else mnote),
            "worked_line": ("tradelab's results by environment, breadth, VIX, state, event ... with 95% intervals"
                            if W is not None else wnote)}
    write_readme(wb, S, info, sheets)
    out = ex / name
    tmp = ex / f".{Path(name).stem}.partial.xlsx"
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    except PermissionError:
        if tmp.exists():
            tmp.unlink()
        raise TLError(f"cannot write {out}: it is open in Excel (close it) or the folder is read-only") from None
    res.update({"file": out, "seconds": round(time.time() - t0, 1), "rows": len(X), "missing": missing,
                "bytes": out.stat().st_size})
    log(f"DONE {out}  ({res['bytes'] / 1e6:,.1f} MB, {res['seconds']:.0f} s): sheets {sheet_name} ({len(X):,} rows), "
        f"Market days ({len(M) if M is not None else 0:,}), What worked ({len(W) if W is not None else 0:,}), Read me")
    return res


# ----------------------------------------------------------------------------------------------------------------
# self-test: a synthetic export with every special case planted, then the journal read back cell by cell
# ----------------------------------------------------------------------------------------------------------------
def _synthetic_export(folder: Path, seed: int = 11) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    days = pd.bdate_range("2024-01-01", periods=30)
    envs = ["broad_bullish", "narrow_bullish", "mixed", "broad_bearish", "high_stress"]
    states = ["S01", "S02", "S03", "S06", "S07", "S08", "S09", "S12", "UNKNOWN", "INSUFFICIENT_DATA"]
    rows = []
    for di, d in enumerate(days[:-6]):
        env = envs[di % 5]
        mkt = float(rng.normal(20, 80))
        for e in ("ENS", "D"):                       # deliberately written ENS first: the journal must reorder
            for rank in (3, 1, 2, 4, 7):             # and shuffled ranks
                k = len(rows)
                net = float(rng.normal(30, 450))
                st = states[k % len(states)]
                rows.append({
                    "trade_id": f"{e}-{d:%Y%m%d}-{rank}", "engine": e, "signal_day": d, "symbol": f"SYM{k % 23:02d}",
                    "rank": rank, "traded": rank <= 3, "both_engines": (k % 9 == 0) and rank <= 3,
                    "other_rank": float(rank) if k % 9 == 0 else np.nan, "repeat": k % 11 == 0,
                    "repeat_of": f"{e}-PREV-{rank}" if k % 11 == 0 else "", "status": "ok", "note": "",
                    "entry_day": (d + pd.offsets.BDay(1)).strftime("%Y-%m-%d"), "entry": 100.0 + k,
                    "entry_gap": float(rng.normal(0, 0.01)), "exit_day": (d + pd.offsets.BDay(5)).strftime("%Y-%m-%d"),
                    "exit_price": (100.0 + k) * (1 + (net + 35) / 1e4), "exit_kind": "close", "sessions_held": 5,
                    "net_bp": net, "mkt_net_bp": mkt, "excess_bp": net - mkt,
                    "mfe": abs(float(rng.normal(0.04, 0.03))), "mae": -abs(float(rng.normal(0.04, 0.03))),
                    "mfe_session": int(rng.integers(1, 6)), "mae_session": int(rng.integers(1, 6)),
                    "touch_up5_session": float(rng.integers(1, 6)) if k % 3 == 0 else np.nan,
                    "touch_dn5_session": float(rng.integers(1, 6)) if k % 4 == 0 else np.nan,
                    **{f"close_ret_s{j}": float(rng.normal(0, 0.03)) for j in range(1, 6)},
                    "pre_state": st, "pre_state_confidence": "UNKNOWN" if st in ("UNKNOWN", "INSUFFICIENT_DATA") else
                    ["HIGH", "MEDIUM", "LOW"][k % 3], "pre_state_phase": "" if st in ("UNKNOWN", "INSUFFICIENT_DATA")
                    else ["EARLY", "MATURE", "DETERIORATING"][k % 3], "pre_state_age": float(1 + k % 17),
                    "pre_state_prev": "S07" if k % 2 else "", "pre_state_pending": "S09" if k % 13 == 0 else "",
                    "pre_state_event": ["", "BREAKOUT:NEW", "BREAKOUT:CONFIRMED", "FAILED_BREAKOUT:FAILED",
                                        "BREAKDOWN:NEW"][k % 5], "pre_state_event_age": float(k % 4),
                    "pre_state_risks": ["", "BIG_DOWN_CANDLE", "GAP_RISK,S05_OVEREXTENDED"][k % 3],
                    "pre_state_notes": ["", "VOL_COMPRESSED"][k % 2],
                    "pre_state_vol_regime": ["NORMAL", "ELEVATED", "EXTREME", "COMPRESSED"][k % 4],
                    "pre_stock_ret": float(rng.normal(0, 0.02)), "pre_state_r20": float(rng.normal(0, 0.1)),
                    "pre_state_rsi14": float(rng.uniform(20, 80)),
                    "pre_home_index_name": ["Smallcap 250", "Midcap 150", "NIFTY 50"][k % 3],
                    "pre_beta_home": float(rng.uniform(0.3, 1.8)), "pre_size_class": "no list",
                    "pre_environment": env, "pre_stress_reason": "VIX at its 97th percentile of 250 sessions"
                    if env == "high_stress" else "", "pre_nifty_ret": float(rng.normal(0, 0.01)),
                    "pre_small_ret": float(rng.normal(0, 0.012)), "pre_nifty_ret20": float(rng.normal(0, 0.04)),
                    "pre_ad_ratio": float(rng.uniform(0.1, 0.9)), "pre_ad_ratio_10": float(rng.uniform(0.3, 0.7)),
                    "pre_vix": float(rng.uniform(10, 25)), "pre_vix_chg5_pct": float(rng.normal(0, 0.15)),
                    "pre_vix_pct250": float(rng.uniform(0, 1)), "post_entry_environment": envs[(di + 1) % 5],
                    "post_hold_home_ret": float(rng.normal(0, 0.02)), "post_hold_nifty_ret": float(rng.normal(0, 0.02)),
                    "post_hold_small_ret": float(rng.normal(0, 0.02)), "post_hold_ad_ratio_avg": float(rng.uniform(0.3, 0.7)),
                    "post_hold_stress_days": float(k % 3), "post_mkt_part_bp": float(rng.normal(0, 150)),
                    "post_stock_part_bp": float(rng.normal(0, 300)),
                    "post_driver": ["market", "stock", "both", "opposite", "flat"][k % 5],
                    "story": "tradelab's own story", "outcome": ""})
    T = pd.DataFrame(rows)
    # planted special cases (all traded): a severe loser, a locked exit, a stuck trade, a net of exactly 0,
    # a net of exactly -500 bp (severe by the rule), and a formula-looking text
    T.loc[0, ["net_bp"]] = -1234.5
    T.loc[1, ["exit_kind", "sessions_held", "net_bp"]] = ["locked_open", 8, -650.0]
    T.loc[2, ["status", "exit_kind", "net_bp", "exit_price", "excess_bp", "note"]] = \
        ["stuck", "stuck_close", np.nan, np.nan, np.nan, "locked for more than 20 sessions"]
    T.loc[5, "net_bp"] = 0.0
    T.loc[6, "net_bp"] = -500.0
    T.loc[7, "pre_stress_reason"] = "=HYPERLINK(\"x\")"
    T["excess_bp"] = np.where(T["status"] == "ok", T["net_bp"] - T["mkt_net_bp"], np.nan)
    folder.mkdir(parents=True, exist_ok=True)
    T.to_parquet(folder / "trades.parquet", index=False)
    M = pd.DataFrame({"date": days, "weekday": [d.strftime("%a") for d in days],
                      "nifty_ret": rng.normal(0, 0.01, len(days)), "mid_ret": rng.normal(0, 0.01, len(days)),
                      "small_ret": rng.normal(0, 0.012, len(days)), "nifty_ret20": rng.normal(0, 0.04, len(days)),
                      "small_ret20": rng.normal(0, 0.05, len(days)), "nifty_vs_ma50": rng.normal(0, 0.03, len(days)),
                      "ad_ratio": rng.uniform(0.2, 0.8, len(days)), "ad_ratio_10": rng.uniform(0.3, 0.7, len(days)),
                      "adv": rng.integers(500, 1500, len(days)), "dec": rng.integers(500, 1500, len(days)),
                      "vix": rng.uniform(10, 25, len(days)), "vix_chg_pct": rng.normal(0, 0.05, len(days)),
                      "vix_chg5_pct": rng.normal(0, 0.15, len(days)), "vix_pct250": rng.uniform(0, 1, len(days)),
                      "environment": [envs[i % 5] for i in range(len(days))],
                      "stress_reason": ["VIX spike" if i % 5 == 4 else "" for i in range(len(days))],
                      "summary": [f"day {i}" for i in range(len(days))]})
    M.to_parquet(folder / "market_days.parquet", index=False)
    B = pd.DataFrame([{"engine": e, "dimension": dim, "timing": tm, "bucket": b, "trades": 10 + i, "share": 0.2,
                       "days": 8 + i, "net_bp_mean": m, "net_bp_lo": m - 50, "net_bp_hi": m + 50,
                       "excess_bp_mean": m / 2, "excess_bp_lo": m / 2 - 80, "excess_bp_hi": m / 2 + 80,
                       "post_mkt_part_bp_mean": 10.0, "post_mkt_part_bp_lo": -5.0, "post_mkt_part_bp_hi": 25.0,
                       "post_stock_part_bp_mean": m - 10, "post_stock_part_bp_lo": m - 60, "post_stock_part_bp_hi": m + 40,
                       "win_rate": 0.5, "severe_rate": 0.15}
                      for e in ("D", "ENS") for i, (dim, tm, b, m) in enumerate([
                          ("what drove the trade", "after", "stock", 120.0),
                          ("environment at the pick", "before", "broad_bullish", 90.0),
                          ("environment at the pick", "before", "high_stress", -120.0),
                          ("stock state at the pick", "before", "S09", 10.0)])])
    B.to_csv(folder / "context_breakdown.csv", index=False)
    (folder / "outputs.sha256").write_text(
        "".join(f"{sha256_file(folder / f)}  {f}\n" for f in ("context_breakdown.csv", "market_days.parquet",
                                                             "trades.parquet")), encoding="utf-8")
    return T


def selftest() -> int:
    _need_openpyxl()
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp(prefix="tl_journal_selftest_"))
    fails = []

    def check(cond, what):
        if not cond:
            fails.append(what)
        return cond
    try:
        ex = tmp / "export" / "TLX_TEST_001"
        T = _synthetic_export(ex)
        before = {p.name: sha256_file(p) for p in ex.iterdir()}
        res = build(ex, log=lambda m: None)
        after = {p.name: sha256_file(p) for p in ex.iterdir() if p.name != DEFAULT_NAME}
        check(before == after, "the export's own files changed")
        check(sorted(p.name for p in ex.iterdir()) == sorted(list(before) + [DEFAULT_NAME]), "extra files written")
        wb = load_workbook(res["file"])
        check(wb.sheetnames == ["Trades", "Market days", "What worked", "Read me"], f"sheets {wb.sheetnames}")
        ws = wb["Trades"]
        heads = [c.value for c in ws[3]]
        H = {h: i for i, h in enumerate(heads)}
        data = [[c for c in row] for row in ws.iter_rows(min_row=4)]
        X = order_rows(T[T["traded"]])
        check(len(data) == len(X) == int(T["traded"].sum()), f"rows {len(data)} vs {len(X)}")
        check(ws.freeze_panes == "G4", f"freeze {ws.freeze_panes}")
        check(ws.auto_filter.ref == f"A3:{ws.cell(3, len(heads)).column_letter}{3 + len(X)}", "filter range")
        # order: newest first, D before ENS, rank 1 first
        keys = [(r[H["Pick day"]].value, r[H["Engine"]].value, r[H["Rank"]].value) for r in data]
        want = sorted(keys, key=lambda k: (-k[0].toordinal(), ENGINE_ORDER.index(k[1]), k[2]))
        check(keys == want, "row order")
        bad_val, bad_col, bad_story = [], [], []
        for row, src in zip(data, X.to_dict("records")):
            v = {h: row[i].value for h, i in H.items()}
            exp_res = result_of(src)
            if v["Symbol"] != src["symbol"] or v["Engine"] != src["engine"] or v["Rank"] != src["rank"]:
                bad_val.append(("identity", src["trade_id"]))
            if pd.Timestamp(v["Pick day"]).date() != pd.Timestamp(src["signal_day"]).date():
                bad_val.append(("day", src["trade_id"]))
            if C.fin(src["net_bp"]):
                if v["Net %"] is None or abs(v["Net %"] - src["net_bp"] / 1e4) > 1e-12:
                    bad_val.append(("net", src["trade_id"]))
            elif v["Net %"] is not None:
                bad_val.append(("net blank", src["trade_id"]))
            if v["Result"] != exp_res:
                bad_val.append(("result", src["trade_id"], v["Result"], exp_res))
            for j in range(1, 6):
                if abs((v[f"Day {j}"] or 0) - src[f"close_ret_s{j}"]) > 1e-12:
                    bad_val.append((f"day {j}", src["trade_id"]))
            if v["Breadth"] is None or abs(v["Breadth"] - src["pre_ad_ratio"]) > 1e-12:
                bad_val.append(("breadth", src["trade_id"]))
            want_notes = ([f"also {'D' if src['engine'] == 'ENS' else 'ENS'} #{int(src['other_rank'])}"]
                          if src["both_engines"] is True or src["both_engines"] == 1 else []) + \
                         ([f"repeat of {src['repeat_of']}"] if src["repeat"] is True or src["repeat"] == 1 else []) + \
                         ([src["note"]] if src["note"] else [])
            if (v["Notes"] or "") != "; ".join(want_notes):
                bad_val.append(("notes", src["trade_id"], v["Notes"], want_notes))
            fill = row[H["Result"]].fill.fgColor.rgb[-6:]
            if fill != RESULT_STYLE[exp_res][0]:
                bad_col.append(("result colour", src["trade_id"], fill))
            fe = row[H["Environment"]].fill.fgColor.rgb[-6:]
            if fe != ENV_STYLE[src["pre_environment"]][0]:
                bad_col.append(("environment colour", src["trade_id"], fe))
            if src["pre_state"] in STATE_STYLE and row[H["State"]].fill.fgColor.rgb[-6:] != STATE_STYLE[src["pre_state"]][0]:
                bad_col.append(("state colour", src["trade_id"]))
            s = v["Story"]
            if exp_res in ("Winner", "Loser", "Severe loser"):
                if f"net {C.pct(src['net_bp'] / 1e4, 2)} after costs - {exp_res.upper()}" not in s:
                    bad_story.append(("net/result", src["trade_id"]))
            elif "stuck" not in s:
                bad_story.append(("stuck", src["trade_id"]))
            if C.ENV_TEXT[src["pre_environment"]] not in s or src["symbol"] not in s:
                bad_story.append(("env/symbol", src["trade_id"]))
            if src["pre_state"] not in ("UNKNOWN", "INSUFFICIENT_DATA") and src["pre_state"] not in s:
                bad_story.append(("state", src["trade_id"]))
        check(not bad_val, f"values {bad_val[:5]}")
        check(not bad_col, f"colours {bad_col[:5]}")
        check(not bad_story, f"story {bad_story[:5]}")
        res_counts = pd.Series([r[H["Result"]].value for r in data]).value_counts().to_dict()
        check(res_counts.get("Stuck") == 1 and res_counts.get("Severe loser", 0) >= 3, f"planted results {res_counts}")
        planted = {r[H["Symbol"]].value: r for r in data}
        r0 = [r for r in data if r[H["Exit"]].value == "locked: sold day 8 open"]
        check(len(r0) == 1 and r0[0][H["Result"]].value == "Severe loser", "locked exit")
        z = [r for r in data if r[H["Net %"]].value is not None and abs(r[H["Net %"]].value) < 1e-15]
        check(len(z) == 1 and z[0][H["Result"]].value == "Loser", "net exactly 0 is a loser (tradelab: winner = net > 0)")
        m5 = [r for r in data if r[H["Net %"]].value is not None and abs(r[H["Net %"]].value + 0.05) < 1e-12]
        check(len(m5) == 1 and m5[0][H["Result"]].value == "Severe loser", "-500 bp is severe")
        fx = [r for r in data if isinstance(r[H["Stress reason"]].value, str) and "HYPERLINK" in r[H["Stress reason"]].value]
        check(len(fx) == 1 and fx[0][H["Stress reason"]].data_type == "s", "formula-looking text kept as text")
        check(len(ws.conditional_formatting) > 10, "conditional formats")
        # every column keeps its own width; the foldable bands carry an outline level
        spec = [(c.head, c.width, bd.fold and len(bd.cols) > 1 and j > 0)
                for bd in trade_bands(False, False, {"D": "ENS", "ENS": "D"}) for j, c in enumerate(bd.cols)]
        dims = [(d.min, d.max, d.width, d.outline_level) for d in ws.column_dimensions.values() if d.min]
        wbad = []
        for k, (hd, w, fold) in enumerate(spec, start=1):
            hit = [x for x in dims if x[0] <= k <= x[1]]
            if not hit or abs((hit[0][2] or 0) - w) > 0.01 or bool(hit[0][3]) != fold:
                wbad.append((hd, w, hit[:1]))
        check(len(spec) == len(heads) and not wbad, f"column widths / outline {wbad[:4]}")
        # the colour keys themselves: every category has its own colour, and the Read me shows the colours used
        for nm, mp in (("result", RESULT_STYLE), ("environment", ENV_STYLE), ("state", STATE_STYLE),
                       ("volatility", VOL_STYLE), ("verdict", VERDICT_STYLE)):
            fl = [v[0] for v in mp.values()]
            check(len(set(fl)) == len(fl), f"{nm} colours not distinct")
        legend_fill = {}
        for row in wb["Read me"].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.fill is not None and c.fill.fgColor is not None and \
                        c.fill.fill_type == "solid":
                    legend_fill[c.value] = c.fill.fgColor.rgb[-6:]
        for k, v in RESULT_STYLE.items():
            check(legend_fill.get(k) == v[0], f"Read me colour for {k}")
        for k, v in ENV_STYLE.items():
            check(legend_fill.get(C.ENV_TEXT[k]) == v[0], f"Read me colour for {k}")
        for k, v in STATE_STYLE.items():
            check(legend_fill.get(state_label(k)) == v[0], f"Read me colour for {k}")
        del planted
        # market days: one row per session, newest first, each day's D average = the source's
        wm = wb["Market days"]
        mh = {c.value: i for i, c in enumerate(wm[3])}
        md = list(wm.iter_rows(min_row=4, values_only=True))
        check(len(md) == 30, f"market rows {len(md)}")
        check([r[mh["Date"]] for r in md] == sorted([r[mh["Date"]] for r in md], reverse=True), "market order")
        G = T[T["traded"] & (T["engine"] == "D")].groupby("signal_day")["net_bp"].mean() / 1e4
        badm = [r[mh["Date"]] for r in md if pd.Timestamp(r[mh["Date"]]) in G.index and
                abs((r[mh["D top-3 net %"]] or 0) - G[pd.Timestamp(r[mh["Date"]])]) > 1e-12]
        check(not badm, f"market D averages {badm[:3]}")
        # what worked: every row, D first, 'before' before 'after', verdicts from the intervals
        ww = wb["What worked"]
        wh = {c.value: i for i, c in enumerate(ww[3])}
        wd = list(ww.iter_rows(min_row=4, values_only=True))
        check(len(wd) == 8, f"what-worked rows {len(wd)}")
        check([r[wh["Engine"]] for r in wd] == ["D"] * 4 + ["ENS"] * 4, "what-worked engine order")
        check([r[wh["Known"]] for r in wd][:4] == ["at the pick"] * 3 + ["afterwards"], "what-worked timing order")
        vv = {(r[wh["Engine"]], r[wh["Group"]]): r[wh["Net vs 0"]] for r in wd}
        check(vv[("D", "Broad bullish")] == "above 0" and vv[("D", "HIGH STRESS")] == "below 0" and
              vv[("D", "S09 Emerging bear")] == "spans 0", f"verdicts {vv}")
        rm = [r[0] for r in wb["Read me"].iter_rows(values_only=True) if r and r[0]]
        check(any("matches the export's outputs.sha256" in str(x) for x in rm), "hash line")
        # --all: every rank, a Traded column, the freeze still right after Net %
        res2 = build(ex, include_all=True, name="ALL.xlsx", log=lambda m: None)
        w2 = load_workbook(res2["file"], read_only=True)
        check(w2.sheetnames[0] == "Picks", "--all sheet name")
        n2 = sum(1 for _ in w2["Picks"].iter_rows(min_row=4))
        check(n2 == len(T), f"--all rows {n2} vs {len(T)}")
        w2.close()
        check(load_workbook(res2["file"])["Picks"].freeze_panes == "H4", "--all freeze")
        (ex / "ALL.xlsx").unlink()
        # absent world: no market days, no breakdown, no state columns -> the journal still builds and says so
        ex2 = tmp / "export" / "TLX_TEST_002"
        ex2.mkdir(parents=True)
        T2 = T.drop(columns=[c for c in T.columns if c.startswith("pre_state")])
        T2["both_engines"] = np.nan                  # blanks must read as 'no', never as True
        T2["repeat"] = np.nan
        T2.to_parquet(ex2 / "trades.parquet", index=False)
        r3 = build(ex2, log=lambda m: None)
        w3 = load_workbook(r3["file"])
        check("not in this export" in str(w3["Market days"]["A1"].value), "absent market days")
        check("not in this export" in str(w3["What worked"]["A1"].value), "absent breakdown")
        notes = " ".join(str(r[0]) for r in w3["Read me"].iter_rows(values_only=True) if r and r[0])
        check("pre_state" in notes and "no outputs.sha256" in notes, "absent columns noted")
        n3 = [c.value for c in w3["Trades"][3]].index("Notes")
        nv = [r[n3] for r in w3["Trades"].iter_rows(min_row=4, values_only=True)]
        check(not any(x and ("also" in x or "repeat" in x) for x in nv), "blank both_engines / repeat read as True")
        # a changed trades.parquet is refused
        T.iloc[:5].to_parquet(ex / "trades.parquet", index=False)
        try:
            build(ex, log=lambda m: None)
            check(False, "a trades.parquet that does not match outputs.sha256 was accepted")
        except TLError:
            pass
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if fails:
        say(f"SELFTEST FAILED ({len(fails)}): " + " | ".join(fails))
        return 1
    say(f"SELFTEST PASSED ({TOOL}, {VERSION}, pandas {pd.__version__}): order, values, colours, story, market days, "
        f"what worked, --all, missing files and columns, the hash refusal")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="tradelab journal: one readable Excel file for an export")
    ap.add_argument("--export", help="an export folder: TLX_YYYYMMDD_NNN (under %%TRADELAB_OUT%%\\export) or its path")
    ap.add_argument("--all", action="store_true", help="also the picks that were not traded (ranks 4-10)")
    ap.add_argument("--name", default=DEFAULT_NAME, help=f"the file name (default {DEFAULT_NAME})")
    ap.add_argument("--selftest", action="store_true", help="check this tool on a synthetic export")
    ap.add_argument("--tradelab-any", action="store_true", help=f"accept a tradelab version other than {NEEDS}")
    a = ap.parse_args(sys.argv[1:] if argv is None else argv)
    try:
        if VERSION != NEEDS and not a.tradelab_any:
            raise TLError(f"{TOOL} was built for {NEEDS}; this tradelab is {VERSION} (pass --tradelab-any to try)")
        if a.selftest:
            return selftest()
        if not a.name.lower().endswith(".xlsx") or any(ch in a.name for ch in "\\/:"):
            raise TLError("--name must be a plain file name ending in .xlsx")
        build(find_export(a.export), include_all=a.all, name=a.name)
        return 0
    except TLError as e:
        say(f"ERROR: {e}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
